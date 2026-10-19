import os
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from docx_helpers import ACCENT, Doc

out_dir = sys.argv[1]
PNG = os.path.join(out_dir, "png")


def png(name):
    return os.path.join(PNG, name + ".png")


doc = Doc()
d = doc.d

# =====================================================================
# Cover
# =====================================================================
for _ in range(5):
    d.add_paragraph()
t = d.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Replikert bankkonto")
r.bold, r.font.size, r.font.color.rgb = True, Pt(34), ACCENT
t = d.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Forklart fra bunnen: oppgaven, systemdesignet og arkitekturen")
r.font.size, r.font.color.rgb = Pt(16), RGBColor(0x44, 0x44, 0x44)
d.add_paragraph()
t = d.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("IN5020/IN9020 Distribuerte systemer · Obligatorisk oppgave 2 · Gruppe 1")
r.font.size = Pt(12)
t = d.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Oktober 2026")
r.font.size, r.font.color.rgb = Pt(12), RGBColor(0x66, 0x66, 0x66)
for _ in range(4):
    d.add_paragraph()
doc.callout("remember", [
    "Dette dokumentet er skrevet for å gi deg **oversikten først**: hva systemet "
    "gjør, hvorfor det er bygget slik, og hvordan delene henger sammen. Les "
    "kapittel 1–5 først. Da har du hele bildet. Kapittel 6–12 går i dybden på hver del.",
    "Dokumentet beskriver **det ferdige programmet**. Timeout (resend og FAIL), "
    "parallell levering og late join beskrives slik de er designet. Sjekklisten i "
    "kapittel 17 viser hva som må stemme med den endelige koden før du presenterer.",
    "Alle figurene finnes også som redigerbare sider i "
    "`IN5020_A2_diagrammer.drawio` (åpnes på diagrams.net eller i draw.io).",
], title="Slik bruker du dokumentet")

doc.page_break()
doc.h1("Innhold", new_page=False)
doc.toc()

# =====================================================================
# 1  The assignment in one minute
# =====================================================================
doc.h1("1  Oppgaven på ett minutt")
doc.p("Vi skal lage **én bankkonto som finnes i flere kopier samtidig**. Hver kopi "
      "kjører i sitt eget program (en **bankserver**, også kalt en **replika**), "
      "gjerne på forskjellige maskiner. Brukere kan gi kommandoer til hvilken som "
      "helst av replikaene: sette inn penger, legge til rente, spørre om saldoen. "
      "Uansett hvor kommandoen gis, skal **alle replikaene ende med nøyaktig samme saldo**.")
doc.p("For å få til det har vi en egen server i midten, **Message Delivery serveren**. "
      "Den har ingen saldo selv. Den eneste jobben dens er å ta imot transaksjoner "
      "fra replikaene, sette dem i **én felles rekkefølge** og sende dem ut til alle. "
      "Alt kommuniserer med **Java RMI** (Remote Method Invocation), altså vanlige "
      "metodekall som går over nettverket.")
doc.h2("Hele oppgaven i fem punkter")
doc.bullets([
    "**Replikering:** kontoen finnes i N kopier. Én kopi kan forsvinne uten at kontoen forsvinner.",
    "**Total rekkefølge:** alle endringer (deposit, addInterest) går via Message Delivery "
    "serveren, som gir dem én global rekkefølge.",
    "**Replikert tilstandsmaskin:** alle replikaer starter med saldo 0 og utfører de samme "
    "endringene i samme rekkefølge, og ender derfor i samme tilstand.",
    "**Medlemskap:** replikaer kan bli med underveis (de får en kopi av tilstanden), "
    "melde seg ut (exit) eller krasje (oppdages med timeout).",
    "**To måter å lese saldoen på:** `getQuickBalance` svarer med en gang fra egen kopi, "
    "`getSyncedBalance` venter til egne transaksjoner er bokført. Den skal lages både "
    "naivt og korrekt, og forskjellen skal forklares.",
])
doc.callout("bank", [
    "Tenk på en bank med **mange filialer** og **ett hovedkontor**. Hver filial har sin egen "
    "kopi av hovedboka for den samme kontoen. Kunder kan levere bilag i hvilken som helst filial.",
    "Filialene bokfører ikke bilagene selv med en gang. De legger dem i en **innkurv**, og en "
    "**kurer** tar innkurven med til hovedkontoret med jevne mellomrom. Hovedkontoret har en "
    "**nummerautomat**: hvert bilag får et løpenummer, og en kopi sendes til **alle** filialer. "
    "Filialene bokfører strengt i løpenummer-rekkefølge. Derfor er alle hovedbøkene like.",
])
doc.h2("Hva som skal leveres")
doc.bullets([
    "Kildekode med gode kommentarer, og to ferdige jar-filer: `MessageDeliveryServer.jar` og `BankServer.jar`.",
    "En README som forklarer hvordan programmet kjøres, forskjellen på naiv og korrekt "
    "`getSyncedBalance`, og hvordan arbeidet ble fordelt i gruppa.",
    "Output-logger fra **alle** replikaene, med tidsstempel på hver linje.",
    "Frist: torsdag 15. oktober 2026 kl. 23:59, i Devilry.",
])

# =====================================================================
# 2  Concepts
# =====================================================================
doc.h1("2  Fem begreper du må kunne")
doc.p("Disse begrepene er grunnmuren. Kan du forklare dem med egne ord og et eksempel, "
      "kan du forklare hele oppgaven.")

doc.h2("2.1  Replikering")
doc.p("Å replikere betyr å ha **flere kopier av de samme dataene** på forskjellige steder. "
      "Det gir to fordeler: systemet tåler at én kopi forsvinner (**feiltoleranse**), og "
      "brukere kan lese fra den nærmeste kopien (**ytelse og tilgjengelighet**).")
doc.p("Prisen er at kopiene må holdes like. Det er hele problemet i denne oppgaven: hvordan "
      "sørger vi for at N kopier av en bankkonto alltid er enige om saldoen?")

doc.h2("2.2  Replikert tilstandsmaskin")
doc.p("En **tilstandsmaskin** har en tilstand og en mengde operasjoner som endrer den. "
      "Bankkontoen vår er en tilstandsmaskin: tilstanden er saldoene per valuta, og "
      "operasjonene er `deposit` og `addInterest`. Operasjonene er **deterministiske**: "
      "samme tilstand + samme operasjon gir alltid samme nye tilstand.")
doc.p("En **replikert tilstandsmaskin** er N kopier av den samme maskinen. Regelen er enkel:")
doc.callout("remember", [
    "Samme starttilstand + samme operasjoner + **samme rekkefølge** = samme sluttilstand.",
    "Alle replikaene starter med saldo 0.0. Hvis vi bare sørger for at alle utfører de samme "
    "transaksjonene i samme rekkefølge, blir de automatisk like. Vi trenger aldri å sammenligne "
    "saldoer eller sende saldoer rundt (unntatt når en ny replika blir med).",
])

doc.h2("2.3  Total rekkefølge (totally ordered broadcast)")
doc.p("En **broadcast** sender en melding til alle i en gruppe. En **totalt ordnet broadcast** "
      "garanterer i tillegg at **alle mottar alle meldinger i nøyaktig samme rekkefølge**, "
      "uansett hvem som sendte dem og når.")
doc.p("Det finnes flere måter å få det til på. Vi bruker den enkleste: en **sequencer**, altså én "
      "sentral prosess som bestemmer rekkefølgen. Det er Message Delivery serveren. Alle "
      "meldinger går via den, den gir dem et løpenummer (meldings-id 0, 1, 2, …), og den sender "
      "dem ut én om gangen. Kapittel 5 viser med et regneeksempel hvorfor rekkefølgen betyr så mye.")

doc.h2("2.4  Sekvensiell konsistens")
doc.p("Oppgaven krever **sekvensiell konsistens** for deposit, addInterest, getHistory, "
      "cleanHistory, memberInfo og saldo-operasjonene. Definisjonen (Lamport, 1979):")
doc.callout("remember", [
    "Resultatet av en kjøring er det samme som om alle operasjonene fra alle prosessene ble "
    "utført i **én bestemt rekkefølge**, og operasjonene fra **hver enkelt prosess** står i den "
    "rekkefølgen i den rekkefølgen prosessen ga dem.",
])
doc.p("Hos oss betyr det:")
doc.bullets([
    "**Én rekkefølge for alle:** Message Delivery serveren lager den globale rekkefølgen, og "
    "alle replikaene utfører i den.",
    "**Hver replikas egen rekkefølge bevares:** `outstanding_collection` er en liste, og den "
    "sendes som en liste. Transaksjonene fra én replika kommer derfor ut i samme rekkefølge som "
    "de ble gitt.",
    "**Lesing:** `getSyncedBalance` (korrekt versjon) er en markør i den samme rekkefølgen, så "
    "svaret passer inn i den globale rekkefølgen. `getQuickBalance` er bevisst svakere: den leser "
    "egen kopi med en gang og kan mangle transaksjoner som ikke er ordnet ennå. Det står i "
    "oppgaveteksten at den skal fungere slik.",
])

doc.h2("2.5  Gruppe og medlemskap")
doc.p("Replikaene for én konto danner en **gruppe** med navnet `<account name>`, for eksempel "
      "`group01`. Hvem som er med i gruppa akkurat nå, kalles **medlemskapet** eller **viewet**. "
      "Message Delivery serveren eier medlemslista og sender den nye lista til alle "
      "(`membershipChanged`) hver gang noen blir med, melder seg ut eller blir erklært død (FAIL).")

# =====================================================================
# 3  Bank analogy
# =====================================================================
doc.h1("3  Bankanalogien")
doc.p("Analogien under er nyttig når du skal forklare systemet muntlig. Hver rad kobler noe "
      "alle forstår fra en vanlig bank til en del av koden.")
doc.figure(png("02_analogi"), "Bankanalogien. Venstre kolonne er banken, høyre kolonne er "
           "den tilsvarende delen av systemet.", width_cm=15.5)
doc.h2("Historien, fortalt som en bank")
doc.steps([
    "En kunde kommer til **filial A** og vil sette inn 100 kroner (`deposit NOK 100`). "
    "Filialen fyller ut et bilag og legger det i **innkurven** (`outstanding_collection`). "
    "Bilaget får filialens eget referansenummer, for eksempel «A 4».",
    "Hvis kunden spør om saldoen **med en gang** (`getQuickBalance`), ser filialen i sin egen "
    "hovedbok. Bilaget ligger fortsatt i innkurven, så de 100 kronene er ikke med ennå.",
    "Hvert 10. sekund tar **kureren** (broadcast-tråden) med seg innkurven til **hovedkontoret** "
    "(Message Delivery serveren).",
    "Hovedkontoret trekker et løpenummer fra **nummerautomaten**, for eksempel melding 7, og "
    "sender en kopi av bilagene til **alle filialene**, også tilbake til A.",
    "Hver filial **bokfører** bilagene i løpenummer-rekkefølge (`receiveTransaction`), "
    "stempler dem som «bokført» (`executedIds`) og sender en **kvittering** (ACK) til hovedkontoret.",
    "Først når hovedkontoret har kvittering fra alle, sender det ut melding 8. Svarer en filial "
    "ikke, ringer hovedkontoret igjen etter 2 sekunder (**resend**). Svarer den fortsatt ikke "
    "etter 5 sekunder, regnes den som stengt (**FAIL**) og fjernes fra lista.",
    "Åpner det en **ny filial**, får den en fotokopi av hovedboka (**state transfer**) før den "
    "begynner å ta imot nye bilag.",
])
doc.callout("say", [
    "«Hver bankserver er en filial med sin egen kopi av hovedboka. Message Delivery serveren er "
    "hovedkontoret med nummerautomaten: den bokfører ingenting selv, men bestemmer rekkefølgen. "
    "Fordi alle filialene bokfører de samme bilagene i samme rekkefølge, blir alle hovedbøkene like.»",
])

# =====================================================================
# 4  Architecture
# =====================================================================
doc.h1("4  Arkitekturen: hele systemet")
doc.p("Figur 2 viser alle delene og alle RMI-kallene mellom dem. Det er den viktigste figuren "
      "i dokumentet. Har du den i hodet, kan du plassere alt annet.")
doc.figure(png("01_system"), "Hele systemet. Øverst Message Delivery serveren, nederst tre "
           "replikaer (A, B, C) og en replika D som blir med senere. Tallene på pilene er "
           "RMI-kallene i forklaringen øverst til venstre.")

doc.h2("4.1  De to programmene")
doc.table(["Program", "Rolle", "Starter med"], [
    ["**MessageDeliveryServer.jar**",
     "Sequencer. Tar imot transaksjonslister, gir hver liste en meldings-id, sender dem til "
     "alle replikaene én om gangen, venter på ACK, håndterer resend og FAIL, eier medlemslista "
     "og gjør state transfer til nye replikaer.",
     "`java -jar jars/MessageDeliveryServer.jar [port]`"],
    ["**BankServer.jar**",
     "Én replika av kontoen. Leser kommandoer, lager transaksjoner, sender dem hvert 10. s, "
     "utfører alle ordnede transaksjoner, svarer på saldo- og historikkspørsmål.",
     "`java [-Dsync=naive] -jar jars/BankServer.jar <server> <konto> <N> <kursfil> [kommandofil]`"],
], [3.6, 8.4, 5.0])

doc.h2("4.2  RMI: hvordan programmene snakker sammen")
doc.p("**Java RMI** lar et program kalle metoder på et objekt som ligger i et annet program, "
      "som om det var lokalt. Det skjer i tre steg:")
doc.steps([
    "Message Delivery serveren lager et **RMI-register** på port 1099 og registrerer seg selv "
    "der under navnet `\"MessageDeliveryServer\"`.",
    "En bankserver slår opp navnet i registeret (`registry.lookup`) og får en **stub**: et "
    "lokalt objekt som ser ut som serveren, men som sender hvert metodekall over nettverket.",
    "Bankserveren kaller `join(accountName, this)`. Den sender altså **seg selv** som argument. "
    "Fordi `BankServer` er et `UnicastRemoteObject`, mottar serveren en stub tilbake til banken. "
    "Nå kan serveren kalle metoder på banken (callbacks).",
])
doc.p("Derfor går RMI **begge veier** (Note 3 i oppgaven): begge programmene er RMI-servere.")
doc.table(["Retning", "Metode", "Når"], [
    ["bank → server", "`join(gruppe, bank)`", "ved oppstart"],
    ["bank → server", "`broadcastTransactions(liste)`", "hvert 10. sekund, hvis innkurven ikke er tom"],
    ["bank → server", "`ack(navn, meldingsId)`", "etter at en melding er utført"],
    ["bank → server", "`leave(gruppe, bank)`", "ved exit"],
    ["server → bank", "`receiveTransaction(meldingsId, liste)`", "for hver melding, til alle"],
    ["server → bank", "`membershipChanged(medlemmer)`", "etter join, leave og FAIL"],
    ["server → bank", "`updateState(BankState)`", "til en ny replika (late join)"],
    ["server → bank", "`getServerName()`", "for å vite navnet på en bank"],
], [3.0, 6.5, 7.5])
doc.p("Alt som sendes over RMI blir **serialisert** (gjort om til bytes). Derfor er "
      "`Transaction` og `BankState` merket `Serializable`.")

doc.h2("4.3  Klassene i koden")
doc.table(["Klasse", "Pakke", "Ansvar"], [
    ["`MessageDeliveryServer`", "message_delivery_server", "Køen, leveringsløkka, ACK-sporing, medlemsliste, join/leave, state transfer"],
    ["`MessageDeliveryServerInterface`", "message_delivery_server", "RMI-metodene banken kan kalle på serveren"],
    ["`BankServer`", "bank_server", "Replikaen: tilstand, transaksjoner, utførelse, saldo, historikk, tråder, oppstart og exit"],
    ["`BankServerInterface`", "bank_server", "RMI-metodene serveren kan kalle på banken"],
    ["`BankServerHelper`", "bank_server", "Tolker én kommandolinje og kaller riktig metode i BankServer"],
    ["`CurrencyConverter`", "bank_server", "Leser kursfila og regner om saldoer mellom valutaer via USD"],
    ["`BankState`", "bank_server", "Øyeblikksbilde av tilstanden som sendes til en ny replika"],
    ["`GetSyncedBalanceDemo`", "bank_server", "Demo som viser at den naive getSyncedBalance kan vente for alltid"],
    ["`Transaction`", "common", "`command` + `uniqueId` (+ tidspunkt for utførelse)"],
    ["`Message`", "common", "Meldings-id + liste med transaksjoner (internt i serveren)"],
    ["`Log`, `Timestamp`", "common", "Tidsstempel foran hver linje som skrives ut"],
], [4.6, 3.9, 8.5])

doc.h2("4.4  Inni en bankserver")
doc.p("Hver bankserver har fire deler som vist i figur 2:")
doc.bullets([
    "**Kommandotråden** leser kommandoer fra terminalen eller fra kommandofila (én linje hvert "
    "0,5–1,5 sekund) og sender dem til `BankServerHelper`.",
    "**Innkurven**, `outstanding_collection`: transaksjoner denne banken har laget, men som "
    "ikke er ordnet ennå.",
    "**Broadcast-tråden** sender en kopi av innkurven til serveren hvert 10. sekund.",
    "**Bokføringen**, `receiveTransaction()`: kalles av serveren med ordnede meldinger. Den "
    "endrer saldoen, øker `order_counter`, legger transaksjonen i `executed_list` og "
    "`executedIds`, og fjerner den fra innkurven hvis det var vår egen.",
])

# =====================================================================
# 5  Why total order
# =====================================================================
doc.h1("5  Hvorfor rekkefølgen er alt")
doc.p("Det er lett å tro at det holder at alle replikaene får alle transaksjonene. Det gjør "
      "det ikke. Figur 3 viser hvorfor med et konkret regneeksempel.")
doc.figure(png("03_total_orden"), "Uten felles rekkefølge ender replikaene med ulik saldo. "
           "Med Message Delivery serveren får alle samme rekkefølge og samme saldo.")
doc.callout("example", [
    "Begge replikaene har 100 USD. Samtidig gir filial A `deposit USD 100` og filial B "
    "`addInterest USD 10`.",
    "- Replika 1 får A først: (100 + 100) × 1,10 = **220 USD**",
    "- Replika 2 får B først: 100 × 1,10 + 100 = **210 USD**",
    "Begge har fått de samme to transaksjonene, men er uenige. Innskudd og rente **kommuterer "
    "ikke**: rekkefølgen endrer svaret.",
])
doc.p("Med Message Delivery serveren blir de to transaksjonene melding #7 og #8. Alle replikaene "
      "utfører #7 før #8, og alle ender på 220 USD. Merk at det ikke betyr noe *hvilken* "
      "rekkefølge serveren velger, bare at **alle bruker den samme**. Det er nettopp det Note 2 "
      "i oppgaven sier: vi trenger ikke ta hensyn til «happened-before» mellom bankservere, "
      "bare til et konsistent view.")
doc.callout("say", [
    "«Fordi rente og innskudd ikke kommuterer, holder det ikke at alle får alle transaksjonene. "
    "De må få dem i samme rekkefølge. Message Delivery serveren er en sequencer som gir denne "
    "rekkefølgen, og da er systemet en replikert tilstandsmaskin.»",
])

# =====================================================================
# 6  Life of a transaction
# =====================================================================
doc.h1("6  En transaksjons livsløp")
doc.p("Figur 4 følger én transaksjon, `deposit NOK 100` gitt på bank A, fra kommandoen skrives "
      "til alle tre replikaene har utført den.")
doc.figure(png("04_transaksjon"), "Sekvensdiagram for én transaksjon. Tidsaksen går nedover.")
doc.steps([
    "**Kommandoen leses.** `BankServerHelper` sjekker argumentene og kaller "
    "`createTransaction(\"deposit NOK 100\")`.",
    "**En Transaction lages.** `command = \"deposit NOK 100\"` og "
    "`uniqueId = \"<servernavn> <outstanding_counter>\"`, for eksempel "
    "`\"BankServer-49478-a3f2 0\"`. Den legges i `outstanding_collection`, og "
    "`outstanding_counter` økes. **Saldoen endres ikke ennå.**",
    "**getQuickBalance nå** gir den gamle saldoen, fordi transaksjonen ikke er utført.",
    "**Broadcast.** Innen 10 sekunder sender broadcast-tråden en kopi av hele innkurven: "
    "`broadcastTransactions([...])`.",
    "**Serveren ordner.** Lista får neste meldings-id (her 7) og legges bakerst i køen. Den "
    "sendes når alle har ACK-et melding 6.",
    "**Levering til alle, parallelt.** Serveren kaller `receiveTransaction(7, liste)` på A, B "
    "og C samtidig.",
    "**Utførelse.** Hver replika går gjennom lista i rekkefølge: hopper over id-er den allerede "
    "har utført, utfører resten (saldo NOK += 100), øker `order_counter`, legger transaksjonen i "
    "`executed_list` og `executedIds`, og fjerner den fra `outstanding_collection` hvis den var "
    "vår egen (bare hos A).",
    "**ACK.** Hver replika kaller `ack(navn, 7)`. Når alle tre er kommet, sender serveren "
    "melding 8.",
])
doc.callout("remember", [
    "Bank A bokfører **ikke** sin egen transaksjon når den lages. Den venter, som alle de andre, "
    "til transaksjonen kommer tilbake fra serveren. Bare slik får A den i samme rekkefølge som B og C.",
])

# =====================================================================
# 7  Message Delivery server
# =====================================================================
doc.h1("7  Message Delivery serveren: ordning, ACK, resend og FAIL")
doc.p("Serveren har én arbeidertråd som kjører `processMessages()` i en løkke. Figur 5 viser "
      "løkka og hva som skjer når en bank ikke svarer.")
doc.figure(png("05_ack_fail"), "Leveringsløkka (venstre) og tidslinjen for én melding der "
           "bank C har krasjet (høyre).")
doc.h2("7.1  Løkka steg for steg")
doc.steps([
    "**Ta neste melding fra køen.** `messageQueue.take()` blokkerer til det finnes en melding. "
    "Køen er FIFO, så meldingene sendes i den rekkefølgen de fikk id.",
    "**Øyeblikksbilde av medlemslista.** `pendingAcks` settes til navnene på alle som er med nå.",
    "**Send til alle samtidig.** `receiveTransaction(id, liste)` kalles parallelt, slik at én "
    "treg bank ikke forsinker leveringen til de andre.",
    "**Vent på ACK.** Hver `ack(navn, id)` fjerner banken fra `pendingAcks`. ACK for en annen "
    "melding enn den som pågår ignoreres.",
    "**Etter 2 sekunder:** banker som ikke har svart, får **samme melding på nytt** (resend).",
    "**Etter 5 sekunder:** banker som fortsatt ikke har svart, markeres som **FAIL**, fjernes "
    "fra medlemslista, og alle de andre får ny medlemsliste.",
    "**Neste melding.** Først nå tas neste melding fra køen.",
])
doc.h2("7.2  Hvorfor vente på alle ACK før neste melding?")
doc.bullets([
    "**Samme rekkefølge overalt.** Ingen replika kan få melding 8 før alle har bekreftet melding 7.",
    "**Ingen replika ligger langt bak.** Forskjellen mellom replikaene er maks én melding.",
    "**Et trygt sted for state transfer.** Mellom to meldinger er alle replikaene i nøyaktig "
    "samme tilstand. Der kan en ny replika få en kopi (kapittel 10).",
    "**Prisen:** gruppa går ikke fortere enn den tregeste replikaen, og i verste fall stopper "
    "den i 5 sekunder før en død replika blir fjernet.",
])
doc.h2("7.3  Hva betyr en ACK?")
doc.p("Banken kaller `ack` på slutten av `receiveTransaction()`, **etter** at alle "
      "transaksjonene i meldingen er utført. En ACK betyr altså «mottatt **og** bokført», ikke "
      "bare «mottatt».")
doc.h2("7.4  Leave er ikke krasj")
doc.p("En bank som avslutter med `exit`, kaller `leave()`. Serveren fjerner den fra "
      "medlemslista **og** fra `pendingAcks` med en gang, så den venter aldri 5 sekunder på en "
      "bank som har sagt fra. En krasjet bank kan ikke si fra, så den oppdages med timeout.")
doc.callout("bank", [
    "Når en filial ringer og sier «vi stenger nå», stryker hovedkontoret den fra lista med en "
    "gang. Når en filial bare slutter å svare, ringer hovedkontoret én gang til, og stryker den "
    "først når den heller ikke da svarer.",
])

# =====================================================================
# 8  Data structures
# =====================================================================
doc.h1("8  Datastrukturene i en replika")
doc.p("Oppgaven bestemmer navnene: `executed_list`, `outstanding_collection`, `order_counter` og "
      "`outstanding_counter`. I tillegg har vi saldoene per valuta og `executedIds`.")
doc.figure(png("06_datastrukturer"), "Hvor en transaksjon befinner seg til hver tid, og "
           "hvordan tellerne og getHistory henger sammen.")
doc.h2("8.1  Oversikt")
doc.table(["Felt", "Hva det er", "Endres når"], [
    ["`balances`", "Saldo per valuta, f.eks. {USD: 324.0, EUR: 252.0, NOK: 328.0}", "en ordnet deposit/addInterest utføres"],
    ["`outstanding_collection`", "Egne transaksjoner som ikke er ordnet ennå (innkurven)", "kommando gis (+), egen transaksjon utføres (−)"],
    ["`executed_list`", "Historikken: utførte transaksjoner siden siste cleanHistory", "utførelse (+), cleanHistory (tømmes)"],
    ["`order_counter`", "Antall utførte transaksjoner totalt = neste ordrenummer", "utførelse (+1), aldri nullstilt"],
    ["`outstanding_counter`", "Teller for å lage id-er til egne deposit/addInterest", "ny transaksjon (+1), aldri nullstilt"],
    ["`executedIds`", "Id-ene til **alle** transaksjoner som noen gang er utført", "utførelse (+), aldri tømt"],
], [4.0, 7.5, 5.5])

doc.h2("8.2  Nøyaktig én gang (exactly-once)")
doc.p("Oppgaven sier at hver transaksjon skal ordnes og utføres **én gang**. Det kan likevel "
      "komme duplikater, av to grunner:")
doc.bullets([
    "Broadcast-tråden sender **hele** innkurven hvert 10. sekund. Har ikke forrige sending "
    "kommet tilbake ennå, sendes de samme transaksjonene en gang til.",
    "Serveren sender en melding **på nytt** etter 2 sekunder uten ACK. Kom den første faktisk "
    "fram, får banken den to ganger.",
])
doc.p("Løsningen er `executedIds`: før en transaksjon utføres, sjekker replikaen om id-en "
      "allerede finnes. Finnes den, **hoppes den over, men ACK sendes likevel**. Fordi alle "
      "replikaene ser meldingene i samme rekkefølge, hopper alle over de samme duplikatene.")
doc.callout("warn", [
    "Hvorfor ikke bare sjekke `executed_list`? Fordi `cleanHistory` tømmer den. Etter en "
    "cleanHistory ville replikaen «glemt» hva den hadde utført, og et duplikat ville blitt "
    "utført en gang til. `executedIds` tømmes aldri.",
])

doc.h2("8.3  getHistory og nummereringen")
doc.p("Oppgaven sier at `executed_list` skal skrives ut med nummer fra "
      "`order_counter − executed_list.size()`. Grunnen er at `cleanHistory` tømmer lista, men ikke "
      "telleren. En transaksjon beholder dermed det samme globale nummeret på alle replikaene, "
      "også etter en cleanHistory. Utestående transaksjoner har ikke noe nummer ennå: de får "
      "det først når de er ordnet.")
doc.code([
    "executed_list",
    "3. 2026-10-14 10:00:01.123 deposit NOK 100  [BankServer-49478-a3f2 0]",
    "4. 2026-10-14 10:00:11.456 addInterest NOK 10  [BankServer-50110-91bc 2]",
    "outstanding_collection",
    "deposit NOK -50  [BankServer-49478-a3f2 1]",
    "addInterest NOK 5  [BankServer-49478-a3f2 2]",
])
doc.p("Tidspunktet er når transaksjonen ble **utført på denne replikaen**. Det er litt "
      "forskjellig fra replika til replika, men ordrenummeret er det samme overalt.")

doc.h2("8.4  checkTxStatus")
doc.bullets([
    "**EXECUTED:** id-en finnes i `executedIds`.",
    "**OUTSTANDING:** transaksjonen ligger i `outstanding_collection`.",
    "**UNKNOWN:** ingen av delene.",
    "`checkTxStatus 4` (bare et tall) betyr «transaksjon nummer 4 laget av **denne** banken». "
    "Full id virker også: `checkTxStatus BankServer-49478-a3f2 4`.",
])

doc.h2("8.5  Valuta")
doc.p("Kontoen har én saldo **per valuta**. Kursfila sier hva én enhet er verdt i USD (USD er "
      "alltid 1):")
doc.code(["EUR 1.17", "NOK 0.10"])
doc.p("For å vise saldoen i en valuta regnes alt om til USD, summeres, og deles på kursen til "
      "målvalutaen:")
doc.callout("example", [
    "Saldo: 100 USD og 100 NOK.",
    "- I USD: 100 × 1 + 100 × 0,10 = **110 USD**",
    "- I NOK: 110 / 0,10 = **1100 NOK**",
    "`addInterest NOK 10` endrer bare NOK-saldoen. `addInterest 10` (uten valuta) gir rente på "
    "alle valutaene.",
])

# =====================================================================
# 9  getSyncedBalance
# =====================================================================
doc.h1("9  getSyncedBalance: naiv og korrekt")
doc.p("Dette er den delen oppgaven ber oss om å **observere og forklare**. Problemet er at "
      "`getQuickBalance` kan svare før egne transaksjoner er utført. `getSyncedBalance` skal "
      "vente til alle transaksjoner denne banken har gitt **før** kommandoen, er utført.")
doc.figure(png("07_getsyncedbalance"), "Samme kommandoer i begge versjonene. Den korrekte "
           "svarer 100.00, den naive 110.00.")
doc.h2("9.1  Naiv: vent på tom innkurv")
doc.p("Den naive versjonen venter til `outstanding_collection` er tom, og leser så saldoen. "
      "Hos oss gjør en egen tråd dette ved å sjekke hvert 200. millisekund.")
doc.bullets([
    "**Feil svar:** Kommer det nye transaksjoner mens den venter, må de også utføres før "
    "innkurven er tom. Svaret tar da med transaksjoner som ble gitt **etter** forespørselen.",
    "**Kan vente for alltid:** Kommer det stadig nye transaksjoner, blir innkurven aldri tom. "
    "Det er «deadlock»-problemet i oppgaveteksten. Mer presist er det **starvation**: ingen "
    "låser hverandre, men betingelsen blir aldri sann.",
])
doc.h2("9.2  Korrekt: en markør i køen")
doc.p("Den korrekte versjonen legger en **markør** (`getSyncedBalance USD`, id "
      "`\"<navn> sync-0\"`) inn i `outstanding_collection`, bak de transaksjonene som allerede "
      "ligger der. Markøren sendes og ordnes som en vanlig transaksjon. Når den kommer tilbake:")
doc.bullets([
    "Svarer **bare banken som laget den**. Andre replikaer ignorerer den.",
    "Den fjernes fra `outstanding_collection`, men **legges ikke i `executed_list`** og øker "
    "ikke `order_counter`.",
    "Svaret er tilstanden akkurat der markøren står i den globale rekkefølgen: alt før er med, "
    "ingenting etter. Svaret kommer alltid etter én broadcast-runde.",
])
doc.h2("9.3  Ikke-blokkerende")
doc.p("I begge versjonene går kommandoløkka videre mens forespørselen venter, og svaret skrives "
      "ut når det er klart. Det er nødvendig for å se forskjellen: bare hvis nye kommandoer kan "
      "komme inn mens vi venter, kan den naive versjonen ta med dem eller vente for alltid. "
      "Versjonen velges ved oppstart med `-Dsync=naive` eller `-Dsync=correct` (standard).")
doc.h2("9.4  Det vi observerte")
doc.p("Eksempelfila kjørt med én replika i hver modus:")
doc.table(["getSyncedBalance", "Korrekt", "Naiv", "Hvorfor forskjellig"], [
    ["1 (før noe innskudd)", "0.00 USD", "0.00 USD", "Ingenting utestående"],
    ["2 (etter deposit USD 100)", "**100.00 USD**", "**110.00 USD**", "Naiv tar med deposit NOK 100, gitt etter forespørselen"],
    ["3 (etter sleep 10)", "110.00 USD", "110.00 USD", "Alt er utført"],
    ["4 (etter addInterest USD 10)", "**120.00 USD**", "**306.00 USD**", "Naiv tar med seks innskudd gitt etter forespørselen"],
    ["5 (i NOK)", "**3177.00 NOK**", "**6516.40 NOK**", "Naiv tar med alt til slutten av fila"],
    ["6 (siste)", "651.64 USD", "651.64 USD", "Ingen nye kommandoer, begge er like"],
], [4.0, 2.8, 2.8, 7.4])
doc.p("De naive svarene avhenger av timing og kan variere fra kjøring til kjøring. De korrekte "
      "er alltid de samme.")
doc.callout("say", [
    "«Den naive versjonen venter på et øyeblikk med tom innkurv. Da tar den med transaksjoner "
    "som kom etter forespørselen, og under jevn last kommer det øyeblikket aldri. Den korrekte "
    "versjonen gjør forespørselen til en del av den totale rekkefølgen, så svaret er tilstanden "
    "akkurat der forespørselen står, og den svarer alltid etter én runde.»",
])

# =====================================================================
# 10  Membership
# =====================================================================
doc.h1("10  Medlemskap: oppstart, ny replika, utmelding og krasj")
doc.p("Gruppa endrer seg over tid. Figur 8 viser de fire situasjonene.")
doc.figure(png("08_medlemskap"), "De fire medlemskapssituasjonene. Hver av dem ender med at "
           "serveren sender ny medlemsliste til alle.")
doc.h2("10.1  Oppstart: vent på N replikaer")
doc.p("Hver bankserver starter med `<number of replicas>` = N. Etter `join` venter den til "
      "medlemslista har minst N navn (`waitForInitialReplicas`), og **først da** starter "
      "broadcast-tråden og kommandoene.")
doc.p("Hvorfor? De opprinnelige replikaene får ingen kopi av tilstanden; de starter alle på 0. "
      "Hvis A sendte en transaksjon før C var med, ville C aldri fått den, og saldoene ville "
      "vært forskjellige for alltid.")
doc.h2("10.2  Ny replika (late join) og state transfer")
doc.p("En replika som starter mens gruppa allerede kjører, ser allerede minst N medlemmer og "
      "starter med en gang. Men den har saldo 0, mens de andre har kjørt en stund. Den må få "
      "en **kopi av tilstanden**:")
doc.steps([
    "Den nye banken kaller `join`.",
    "Serveren venter til meldingen som pågår er ACK-et av alle. Da er alle replikaene i "
    "nøyaktig samme tilstand, og ingen melding er halvveis levert.",
    "Serveren henter et øyeblikksbilde fra en eksisterende replika: `BankState` med saldoene, "
    "`executed_list`, `order_counter` og `executedIds`.",
    "Serveren kaller `updateState(state)` på den nye banken, som kopierer alt inn.",
    "Den nye banken legges i medlemslista, alle får `membershipChanged`, og fra neste melding "
    "får den alt som de andre.",
])
doc.callout("warn", [
    "Tilstanden og meldingene etterpå må **henge sammen uten hull eller overlapp**. Tas kopien "
    "midt i en melding, kan den nye replikaen mangle en transaksjon eller få den to ganger. "
    "Derfor skjer state transfer mellom to meldinger. `executedIds` må være med, ellers kan den "
    "nye replikaen utføre et duplikat som de andre hopper over.",
])
doc.h2("10.3  Utmelding (exit)")
doc.steps([
    "Send resten av innkurven, slik at de andre utfører transaksjonene selv om vi går.",
    "Vent (maks 30 s) på `getSyncedBalance`-svar som ikke er kommet ennå.",
    "`leave()`: serveren fjerner oss straks og sender ny medlemsliste.",
    "`System.exit(0)`. RMI-trådene ville ellers holdt programmet i live.",
])
doc.h2("10.4  Krasj")
doc.p("En bank som krasjer (kill, Ctrl+C, maskinen går ned), svarer ikke på neste melding. "
      "Serveren sender på nytt etter 2 sekunder og erklærer den FAIL etter 5 sekunder. Den "
      "fjernes, de andre får ny medlemsliste, og resten av gruppa fortsetter med lik tilstand.")
doc.p("Transaksjoner som lå i innkurven til den krasjede banken og **aldri ble sendt**, går tapt. "
      "Kunden fikk heller aldri bekreftet dem. Transaksjoner som **allerede var sendt**, utføres "
      "av alle de andre.")

# =====================================================================
# 11  Threads
# =====================================================================
doc.h1("11  Tråder og låser")
doc.p("En bankserver gjør flere ting samtidig. Figur 9 viser trådene og den delte tilstanden.")
doc.figure(png("09_traader"), "Trådene i en bankserver. Alle som leser eller endrer den delte "
           "tilstanden, gjør det med synchronized.")
doc.table(["Tråd", "Gjør", "Lås"], [
    ["Kommandotråd (main)", "Leser kommandoer, lager transaksjoner, leser historikk og saldo", "`synchronized`"],
    ["Broadcast-tråd", "Kopierer innkurven og sender den hvert 10. s", "`synchronized` mens den kopierer"],
    ["RMI-tråd: receiveTransaction", "Utfører transaksjoner, svarer egne markører, sender ACK", "`synchronized`"],
    ["RMI-tråd: membershipChanged", "Lagrer ny medlemsliste, vekker oppstarten", "egen `membershipLock`"],
    ["RMI-tråd: updateState", "Kopierer inn tilstanden ved late join", "`synchronized`"],
    ["Naiv getSyncedBalance", "Sjekker hvert 200 ms om innkurven er tom (bare -Dsync=naive)", "`synchronized` mens den sjekker"],
], [4.6, 8.0, 4.4])
doc.bullets([
    "**Hvorfor lås:** uten låsen kan `receiveTransaction` fjerne en transaksjon fra innkurven "
    "mens broadcast-tråden kopierer den. Resultatet blir feil, eller "
    "`ConcurrentModificationException`.",
    "**Hvorfor vente uten lås:** `getSyncedBalance` venter på svaret **uten** å holde låsen. "
    "Holdt den låsen, kunne `receiveTransaction` aldri komme inn og levere svaret: vranglås.",
    "**Hvorfor egen lås for medlemskap:** oppstarten venter på medlemslista. Med en egen lås "
    "blokkeres ikke en medlemsoppdatering av en transaksjon som utføres.",
])

# =====================================================================
# 12  Running and testing
# =====================================================================
doc.h1("12  Kjøring og testing")
doc.figure(png("10_kjoring"), "Testscenariet med seks terminaler. Tidspunktene er fra "
           "fjorårets README.txt og må sjekkes mot årets filer.")
doc.h2("12.1  Bygg")
doc.code(["./scripts/build.sh          # macOS / Linux", "scripts\\build.bat           # Windows"])
doc.p("Skriptet kompilerer med `--release 11`, så jar-filene kjører på Java 11 eller nyere.")
doc.h2("12.2  Kjør")
doc.code([
    "# Terminal 1",
    "java -jar jars/MessageDeliveryServer.jar 1099",
    "",
    "# Terminal 2, 3, 4 (ved 0 s)",
    "java -jar jars/BankServer.jar localhost:1099 group01 3 input/final/TradingRate.txt input/final/Rep1.txt",
    "",
    "# Terminal 5 (ved 5 s) og 6 (ved 15 s): samme, med Rep2.txt og Rep3.txt",
    "",
    "# Naiv versjon: java -Dsync=naive -jar jars/BankServer.jar ...",
])
doc.p("Legg til `2>&1 | tee output/final/<navn>.txt` på slutten for å lagre loggen samtidig "
      "som den vises.")
doc.h2("12.3  Hva vi sjekker")
doc.bullets([
    "Alle replikaene som levde til slutten, har **samme sluttsaldo**.",
    "Alle replikaene har **samme transaksjon på hvert ordrenummer** (sammenlign "
    "«Received transaction»-linjene uten tidsstempel med `diff`).",
    "Naiv og korrekt modus gir **forskjellige** svar på noen `getSyncedBalance`, og forskjellen "
    "kan forklares.",
    "En replika som blir drept midt i kjøringen, blir FAIL etter 5 sekunder, og resten fortsetter.",
])

# =====================================================================
# 13  Commands
# =====================================================================
doc.h1("13  Alle kommandoene")
doc.table(["Kommando", "Hva skjer", "Går via serveren?"], [
    ["`deposit <Cur> <beløp>`", "Ny transaksjon i innkurven. Saldoen økes på alle replikaer når den er ordnet.", "Ja"],
    ["`addInterest [<Cur>] <prosent>`", "Som deposit. Saldo × (1 + prosent/100), for én valuta eller alle.", "Ja"],
    ["`getQuickBalance [<Cur>]`", "Leser egen kopi med en gang, omregnet til valutaen (standard USD).", "Nei"],
    ["`getSyncedBalance [<Cur>]`", "Svarer når egne tidligere transaksjoner er utført (kapittel 9).", "Ja (korrekt) / nei (naiv)"],
    ["`getHistory`", "Skriver ut executed_list med nummer og tidspunkt, og outstanding_collection.", "Nei"],
    ["`checkTxStatus <id>`", "EXECUTED, OUTSTANDING eller UNKNOWN.", "Nei"],
    ["`cleanHistory`", "Tømmer executed_list. Tellerne og executedIds beholdes.", "Nei"],
    ["`memberInfo`", "Skriver ut medlemmene i gruppa nå.", "Nei"],
    ["`sleep <sekunder>`", "Venter. Nyttig i kommandofiler.", "Nei"],
    ["`exit`", "Sender resten av innkurven, melder seg ut, avslutter.", "leave()"],
], [4.6, 9.4, 3.0])

# =====================================================================
# 14  Design decisions
# =====================================================================
doc.h1("14  Designvalg og begrunnelser")
doc.p("Dette er valg vi har tatt der oppgaven ikke bestemmer svaret. Det er nyttig å kunne "
      "begrunne dem.")
doc.table(["Valg", "Begrunnelse"], [
    ["`executedIds` i tillegg til `executed_list`", "Duplikatsjekken må virke også etter cleanHistory."],
    ["ACK sendes også når alle transaksjonene var duplikater", "Ellers ville serveren vente for alltid, eller feilaktig erklære banken død."],
    ["Vente på N replikaer før start", "De opprinnelige replikaene får ingen state transfer, så ingen kan sende før alle er med."],
    ["`getSyncedBalance` blokkerer ikke kommandoløkka", "Bare slik kan forskjellen mellom naiv og korrekt observeres, og tidsplanen i kommandofilene holdes."],
    ["Markørene får egne id-er (`sync-0`, `sync-1`, …)", "`outstanding_counter` skal bare nummerere deposit og addInterest, slik at `checkTxStatus N` treffer riktig transaksjon."],
    ["`-Dsync=naive|correct` i stedet for et ekstra argument", "Argumentlista er bestemt av oppgaven, og et sjette argument ville kollidert med den valgfrie kommandofila."],
    ["Valuta mangler → USD", "USD er basisvalutaen i kursfila."],
    ["Kommandoen lagres slik den ble skrevet", "Som i oppgaveteksten («deposit NOK 500»). Alle replikaer tolker den samme teksten likt."],
    ["Servernavn = PID + tilfeldig suffiks", "Navnet er en del av hver transaksjons-id. PID alene er bare unikt på én maskin."],
    ["`leave()` ved exit", "En planlagt utmelding skal ikke vente 5 sekunder og ende som FAIL."],
    ["Tidsstempel på hver linje via `Log.install()`", "Kravet gjelder hver linje. Én endring dekker all output, også feilmeldinger."],
    ["Saldo skrives med to desimaler", "Bare ved utskrift. Den interne verdien er uendret, så alle replikaer regner likt."],
    ["`--release 11`", "Jar-filene skal kjøre hos sensor, som kan ha en eldre Java."],
], [6.0, 11.0])

# =====================================================================
# 15  Q&A
# =====================================================================
doc.h1("15  Spørsmål foreleserne kan stille")
qa = [
    ("Hvorfor trenger vi Message Delivery serveren? Kunne ikke bankene sendt direkte til hverandre?",
     "De kunne, men da kan meldingene komme fram i forskjellig rekkefølge hos forskjellige "
     "replikaer, og saldoene blir ulike (kapittel 5). Uten en sentral server trenger man en "
     "distribuert ordningsalgoritme, for eksempel med Lamport-klokker. En sequencer er den "
     "enkleste måten å få total rekkefølge på."),
    ("Er ikke serveren et single point of failure?",
     "Jo. Dør Message Delivery serveren, stopper hele systemet. Det er en kjent svakhet ved "
     "sequencer-løsningen, og oppgaven ber ikke om å løse den. I et ekte system ville man "
     "replikert sequenceren med en konsensusprotokoll som Paxos eller Raft. Den er også en "
     "flaskehals, siden all trafikk går gjennom den."),
    ("Hva betyr sekvensiell konsistens i systemet deres?",
     "At resultatet er som om alle operasjonene ble utført i én rekkefølge, der hver replikas "
     "egne operasjoner står i den rekkefølgen den ga dem. Serveren lager den ene rekkefølgen, "
     "og innkurven sendes som en liste, så hver replikas rekkefølge bevares."),
    ("Er getQuickBalance feil?",
     "Nei, den er bevisst svakere. Den leser egen kopi med en gang, uten kommunikasjon, og kan "
     "mangle transaksjoner som ikke er ordnet ennå. Den er rask. getSyncedBalance gir sterkere "
     "konsistens, men tar lengre tid."),
    ("Kan den korrekte getSyncedBalance også henge?",
     "Bare hvis serveren eller nettverket henger. Markøren blir ordnet i neste broadcast-runde "
     "uansett hvor mye trafikk det er. En død replika blokkerer heller ikke, fordi den blir "
     "fjernet etter 5 sekunder."),
    ("Hva skjer hvis den samme transaksjonen kommer to ganger?",
     "Replikaen ser id-en i executedIds, hopper over transaksjonen og sender ACK likevel. Alle "
     "replikaene hopper over de samme duplikatene, fordi de ser meldingene i samme rekkefølge."),
    ("Hvorfor sender dere transaksjoner og ikke bare den nye saldoen?",
     "Oppgaven krever en replikert tilstandsmaskin: operasjonene sendes, og hver replika "
     "utfører dem. Det gir også historikken (getHistory) og ordrenumrene. Å sende saldoen ville "
     "dessuten krevd at noen bestemte hvilken saldo som var riktig ved samtidige endringer."),
    ("Hvorfor er replikaene like når tidsstemplene i loggene er forskjellige?",
     "Tidsstemplene er lokale: når hver replika utførte transaksjonen. Det som må være likt, er "
     "rekkefølgen og ordrenummeret, og det er det."),
    ("Hvordan blir en ny replika lik de andre?",
     "Med state transfer: serveren venter til meldingen som pågår er ferdig, henter en kopi av "
     "tilstanden fra en eksisterende replika, gir den til den nye med updateState, og legger "
     "den inn i medlemslista før neste melding."),
    ("Hva skjer hvis en replika bare er treg, ikke død?",
     "Svarer den ikke innen 5 sekunder, blir den erklært FAIL og fjernet, selv om den lever. "
     "Den får da ikke flere meldinger og vil etter hvert ha feil tilstand. Den måtte startet på "
     "nytt og blitt med som en ny replika. Det er den klassiske avveiningen ved timeout-basert "
     "feildeteksjon: man kan ikke skille en treg prosess fra en død."),
    ("Hvorfor RMI begge veier?",
     "Serveren må kunne dytte meldinger til bankene (receiveTransaction) i stedet for at bankene "
     "spør hele tiden. Banken sender seg selv som argument i join, og serveren får en stub "
     "tilbake til banken."),
    ("Hva er forskjellen på executed_list og executedIds?",
     "executed_list er historikken som vises til brukeren og tømmes av cleanHistory. "
     "executedIds brukes internt for å garantere at hver transaksjon bare utføres én gang, og "
     "tømmes aldri."),
    ("Hvorfor sendes innkurven bare hvert 10. sekund?",
     "Det står i oppgaven. Det samler flere transaksjoner i én melding og gir færre meldinger "
     "og ACK-runder. Ulempen er høyere forsinkelse: en transaksjon kan vente opptil 10 sekunder "
     "før den i det hele tatt blir sendt."),
]
for q, a in qa:
    para = d.add_paragraph()
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.keep_with_next = True
    r = para.add_run("Spørsmål: " + q)
    r.bold = True
    r.font.color.rgb = ACCENT
    doc.p("**Svar:** " + a)

# =====================================================================
# 16  Glossary
# =====================================================================
doc.h1("16  Ordliste")
doc.table(["Begrep", "Betydning"], [
    ["Replika", "Én kopi av kontoen, dvs. én kjørende bankserver."],
    ["Sequencer", "Den sentrale prosessen som bestemmer rekkefølgen. Her: Message Delivery serveren."],
    ["Total rekkefølge", "Alle mottar alle meldinger i nøyaktig samme rekkefølge."],
    ["Replikert tilstandsmaskin", "Samme start + samme operasjoner i samme rekkefølge = samme tilstand på alle kopier."],
    ["Sekvensiell konsistens", "Resultatet er som om alle operasjoner ble utført i én rekkefølge som respekterer hver prosess sin rekkefølge."],
    ["ACK", "Kvittering: «meldingen er mottatt og utført»."],
    ["Resend", "Serveren sender samme melding på nytt etter 2 s uten ACK."],
    ["FAIL", "Serveren erklærer en bank død etter 5 s uten ACK og fjerner den."],
    ["Exactly-once", "Hver transaksjon utføres nøyaktig én gang, selv om den mottas flere ganger."],
    ["Idempotent", "Å gjøre noe to ganger har samme effekt som én gang (duplikatsjekken gjør mottaket idempotent)."],
    ["View / medlemskap", "Hvem som er med i gruppa akkurat nå."],
    ["State transfer", "En ny replika får en kopi av tilstanden til de andre."],
    ["Late join", "En replika som blir med etter at gruppa har startet."],
    ["RMI", "Remote Method Invocation: metodekall over nettverket i Java."],
    ["Stub", "Lokalt objekt som videresender metodekall til et objekt i et annet program."],
    ["RMI-register", "Navnetjeneste der serveren registrerer seg, og der bankene slår den opp."],
    ["Serialisering", "Å gjøre et objekt om til bytes for å sende det over nettverket."],
    ["Markør", "getSyncedBalance-forespørselen som ordnes sammen med transaksjonene."],
    ["Starvation", "En betingelse blir aldri sann, så noe venter for alltid (den naive getSyncedBalance)."],
    ["Vranglås (deadlock)", "To tråder venter på hverandre, og ingen kommer videre."],
    ["Daemon-tråd", "En tråd som ikke holder programmet i live når main er ferdig."],
], [4.5, 12.5])

# =====================================================================
# 17  Checklist
# =====================================================================
doc.h1("17  Sjekkliste før presentasjonen")
doc.p("Dokumentet beskriver det ferdige designet. Disse punktene var ikke ferdig implementert "
      "da dokumentet ble skrevet (11. oktober). Sjekk at koden gjør det som står her, eller "
      "juster forklaringen:")
doc.bullets([
    "**Parallell levering:** `processMessages()` sender `receiveTransaction` til alle bankene "
    "samtidig (ikke én og én).",
    "**Resend etter 2 s og FAIL etter 5 s**, med `notifyMembershipChanged()` etter FAIL.",
    "**Late join:** serveren henter `BankState` fra en eksisterende replika mellom to meldinger "
    "og kaller `updateState` på den nye. `BankState` inneholder også `executedIds`.",
    "**`updateState()`** i `BankServer` kopierer saldoene, `executed_list`, `order_counter` og "
    "`executedIds`.",
    "**Testscenariet** stemmer med årets `README.txt` (tidspunktene i figur 10 er fra i fjor).",
    "**Tallene i kapittel 9.4** er fra eksempelfila. Oppdater dem gjerne med resultatene fra "
    "årets filer.",
])
doc.callout("say", [
    "Øv på å forklare **figur 2** (hele systemet) og **figur 3** (hvorfor rekkefølgen er alt) "
    "uten å se på teksten. Kan du de to, kan du svare på nesten alt annet ut fra dem.",
], title="Tips")

doc.save(os.path.join(out_dir, "IN5020_A2_forklaring.docx"))
print("saved")
