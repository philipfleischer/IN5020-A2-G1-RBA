import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)

CALLOUT = {
    "bank": ("FFE6CC", "D79B00", "Bankanalogi"),
    "remember": ("DAE8FC", "6C8EBF", "Husk"),
    "example": ("FFF2CC", "D6B656", "Eksempel"),
    "warn": ("F8CECC", "B85450", "Viktig"),
    "say": ("D5E8D4", "82B366", "Slik kan du si det"),
}


def _shade(el, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


def _cell_borders(cell, color, size=12):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement("w:" + side)
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:color"), color)
        borders.append(b)
    tcPr.append(borders)


def _cell_margins(cell, cm=0.25):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        m = OxmlElement("w:" + side)
        m.set(qn("w:w"), str(int(cm * 567)))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tcPr.append(mar)


class Doc:
    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.orientation = WD_ORIENT.PORTRAIT
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.0))
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        self.width = Cm(17)

        st = self.d.styles
        st["Normal"].font.name = BODY_FONT
        st["Normal"].font.size = Pt(11)
        st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        st["Normal"].paragraph_format.space_after = Pt(6)
        st["Normal"].paragraph_format.line_spacing = 1.12
        for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
            s = st[name]
            s.font.name = BODY_FONT
            s.font.size = Pt(size)
            s.font.color.rgb = ACCENT
            rf = s.element.rPr.rFonts
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), BODY_FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            s.paragraph_format.space_before = Pt(16 if name == "Heading 1" else 12)
            s.paragraph_format.space_after = Pt(6)
            s.paragraph_format.keep_with_next = True
        self.fig = 0
        self._footer()

    # ------------------------------------------------------------ basics
    def _footer(self):
        p = self.d.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("IN5020 Oblig 2 · Replikert bankkonto · side ")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        self._field(p, "PAGE", "1", size=9)

    @staticmethod
    def _field(p, instr, placeholder, size=None):
        r1 = p.add_run()
        f1 = OxmlElement("w:fldChar")
        f1.set(qn("w:fldCharType"), "begin")
        r1._r.append(f1)
        r2 = p.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = " %s " % instr
        r2._r.append(it)
        r3 = p.add_run()
        f3 = OxmlElement("w:fldChar")
        f3.set(qn("w:fldCharType"), "separate")
        r3._r.append(f3)
        r4 = p.add_run(placeholder)
        if size:
            r4.font.size = Pt(size)
        r5 = p.add_run()
        f5 = OxmlElement("w:fldChar")
        f5.set(qn("w:fldCharType"), "end")
        r5._r.append(f5)

    def _runs(self, p, text, size=None, color=None, italic=False):
        """**bold** and `code` inside a paragraph."""
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`"):
                r = p.add_run(part[1:-1])
                r.font.name = CODE_FONT
                r._r.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
                r.font.size = Pt((size or 11) - 1)
                r.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
            else:
                r = p.add_run(part)
            if size and not part.startswith("`"):
                r.font.size = Pt(size)
            if color and not part.startswith("`"):
                r.font.color.rgb = color
            if italic:
                r.italic = True
        return p

    def h1(self, t, new_page=True):
        if new_page:
            self.page_break()
        return self.d.add_heading(t, 1)

    def h2(self, t):
        return self.d.add_heading(t, 2)

    def h3(self, t):
        return self.d.add_heading(t, 3)

    def p(self, text, **kw):
        return self._runs(self.d.add_paragraph(), text, **kw)

    def bullets(self, items, style="List Bullet"):
        for it in items:
            para = self.d.add_paragraph(style=style)
            para.paragraph_format.space_after = Pt(3)
            self._runs(para, it)

    def steps(self, items):
        """Numbered steps without Word list numbering (which never restarts)."""
        for i, it in enumerate(items, 1):
            para = self.d.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.9)
            para.paragraph_format.first_line_indent = Cm(-0.9)
            para.paragraph_format.space_after = Pt(4)
            r = para.add_run("%d.\t" % i)
            r.bold = True
            r.font.color.rgb = ACCENT
            para.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            self._runs(para, it)

    def code(self, lines):
        for i, line in enumerate(lines):
            para = self.d.add_paragraph()
            pf = para.paragraph_format
            pf.space_after = Pt(0) if i < len(lines) - 1 else Pt(8)
            pf.space_before = Pt(0)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.3)
            _shade(para._p.get_or_add_pPr(), "F3F3F3")
            r = para.add_run(line if line else " ")
            r.font.name = CODE_FONT
            r._r.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
            r.font.size = Pt(9)

    def callout(self, kind, paragraphs, title=None):
        fill, border, default = CALLOUT[kind]
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.width = self.width
        _shade(cell._tc.get_or_add_tcPr(), fill)
        _cell_borders(cell, border)
        _cell_margins(cell)
        first = cell.paragraphs[0]
        r = first.add_run(title or default)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(border)
        for txt in paragraphs:
            para = cell.add_paragraph()
            para.paragraph_format.space_after = Pt(3)
            if txt.startswith("- "):
                para.paragraph_format.left_indent = Cm(0.5)
                para.paragraph_format.first_line_indent = Cm(-0.35)
                txt = "•  " + txt[2:]
            self._runs(para, txt, size=10.5)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, header, rows, widths, size=10):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = self.d.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.width = Cm(widths[i])
            _shade(c._tc.get_or_add_tcPr(), "1F4E79")
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # repeat header row on new pages
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                c = cells[i]
                c.width = Cm(widths[i])
                if ri % 2 == 1:
                    _shade(c._tc.get_or_add_tcPr(), "F2F6FB")
                para = c.paragraphs[0]
                para.paragraph_format.space_after = Pt(2)
                for k, line in enumerate(val.split("\n")):
                    if k:
                        para = c.add_paragraph()
                        para.paragraph_format.space_after = Pt(2)
                    self._runs(para, line, size=size)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figure(self, path, caption, width_cm=17):
        self.fig += 1
        para = self.d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.space_after = Pt(2)
        para.add_run().add_picture(path, width=Cm(width_cm))
        cap = self.d.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._runs(cap, "**Figur %d.** %s" % (self.fig, caption), size=9.5,
                   color=RGBColor(0x44, 0x44, 0x44))
        return self.fig

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self):
        para = self.d.add_paragraph()
        self._field(para, 'TOC \\o "1-2" \\h \\z \\u',
                    "Innholdsfortegnelsen fylles inn når Word oppdaterer feltene "
                    "(høyreklikk her → Oppdater felt).")
        settings = self.d.settings.element
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)

    def save(self, path):
        self.d.save(path)
